from flask import Flask, request, jsonify, send_file
from flask_cors import CORS

import joblib
import os
import time
import requests
from io import BytesIO

from pypdf import PdfReader
from docx import Document

from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4


# =========================================================
# PATHS
# =========================================================

BASE_DIR = os.path.dirname(
    os.path.abspath(__file__)
)

FRONTEND_DIR = os.path.join(
    BASE_DIR,
    "frontend"
)


# =========================================================
# FLASK APP
# =========================================================

app = Flask(
    __name__,
    static_folder=FRONTEND_DIR,
    static_url_path=""
)

CORS(app)


# =========================================================
# MODEL PATHS
# =========================================================

MODEL_PATH = os.path.join(
    BASE_DIR,
    "model.pkl"
)

VECTORIZER_PATH = os.path.join(
    BASE_DIR,
    "vectorizer.pkl"
)


# =========================================================
# LOAD MODEL
# =========================================================

print("Loading language detection model...")

model = joblib.load(
    MODEL_PATH
)

vectorizer = joblib.load(
    VECTORIZER_PATH
)

print("MODEL LOADED SUCCESSFULLY")


# =========================================================
# LANGUAGE CODES
# =========================================================

LANGUAGE_CODES = {

    "English": "en",
    "Hindi": "hi",
    "Telugu": "te",
    "Tamil": "ta",
    "Kannada": "kn",
    "Malayalam": "ml",

    "Spanish": "es",
    "French": "fr",
    "German": "de",
    "Portuguese": "pt",
    "Italian": "it",
    "Russian": "ru",
    "Dutch": "nl",

    "Arabic": "ar",
    "Turkish": "tr",
    "Danish": "da",
    "Greek": "el",
    "Swedish": "sv"
}


# =========================================================
# HOME
# =========================================================
# This now opens the FRONTEND instead of returning JSON.
# Therefore frontend + backend use ONE public URL.
# =========================================================

@app.route("/")
def home():

    return app.send_static_file(
        "index.html"
    )


# =========================================================
# HEALTH
# =========================================================

@app.route("/health")
def health():

    return jsonify({

        "success": True,

        "status": "healthy",

        "message":
            "Language Detection Application is running"

    })


# =========================================================
# TEXT LANGUAGE DETECTION
# =========================================================

@app.route(
    "/predict",
    methods=["POST"]
)
def predict():

    try:

        data = request.get_json()

        if not data:

            return jsonify({

                "success": False,

                "message":
                    "No JSON data received."

            }), 400


        text = data.get(
            "text",
            ""
        ).strip()


        if not text:

            return jsonify({

                "success": False,

                "message":
                    "Please enter text."

            }), 400


        # Convert text into ML features

        features = vectorizer.transform(
            [text]
        )


        # Predict language

        prediction = model.predict(
            features
        )[0]


        # Calculate confidence

        if hasattr(
            model,
            "predict_proba"
        ):

            probabilities = (
                model.predict_proba(
                    features
                )[0]
            )

            confidence = (
                max(probabilities) * 100
            )

        else:

            confidence = 0


        return jsonify({

            "success": True,

            "language":
                str(prediction),

            "confidence":
                round(
                    confidence,
                    2
                ),

            "text":
                text

        })


    except Exception as error:

        print(
            "Prediction error:",
            error
        )


        return jsonify({

            "success": False,

            "message":
                str(error)

        }), 500


# =========================================================
# FILE TEXT EXTRACTION
# =========================================================

def extract_text_from_file(file):

    filename = (
        file.filename.lower()
    )


    # -----------------------------------------------------
    # TXT
    # -----------------------------------------------------

    if filename.endswith(".txt"):

        content = file.read()


        try:

            return content.decode(
                "utf-8"
            )

        except UnicodeDecodeError:

            return content.decode(
                "latin-1"
            )


    # -----------------------------------------------------
    # PDF
    # -----------------------------------------------------

    elif filename.endswith(".pdf"):

        reader = PdfReader(
            file
        )

        pages = []


        for page in reader.pages:

            page_text = (
                page.extract_text()
            )


            if page_text:

                pages.append(
                    page_text
                )


        return "\n".join(
            pages
        )


    # -----------------------------------------------------
    # DOCX
    # -----------------------------------------------------

    elif filename.endswith(".docx"):

        document = Document(
            file
        )

        paragraphs = []


        for paragraph in (
            document.paragraphs
        ):

            text = (
                paragraph.text.strip()
            )


            if text:

                paragraphs.append(
                    text
                )


        return "\n".join(
            paragraphs
        )


    # -----------------------------------------------------
    # INVALID FILE
    # -----------------------------------------------------

    else:

        raise ValueError(

            "Only PDF, DOCX and TXT "
            "files are supported."

        )


# =========================================================
# FILE UPLOAD + LANGUAGE DETECTION
# =========================================================

@app.route(
    "/upload",
    methods=["POST"]
)
def upload_file():

    try:

        if "file" not in request.files:

            return jsonify({

                "success": False,

                "message":
                    "No file uploaded."

            }), 400


        file = request.files[
            "file"
        ]


        if not file.filename:

            return jsonify({

                "success": False,

                "message":
                    "Please select a file."

            }), 400


        # Extract text

        extracted_text = (
            extract_text_from_file(
                file
            )
        )


        extracted_text = (
            extracted_text.strip()
        )


        if not extracted_text:

            return jsonify({

                "success": False,

                "message":
                    "No readable text found "
                    "in the file."

            }), 400


        # Use first 10000 characters

        detection_text = (
            extracted_text[:10000]
        )


        # Convert text to features

        features = (
            vectorizer.transform(
                [detection_text]
            )
        )


        # Predict language

        prediction = (
            model.predict(
                features
            )[0]
        )


        # Confidence

        if hasattr(
            model,
            "predict_proba"
        ):

            probabilities = (
                model.predict_proba(
                    features
                )[0]
            )

            confidence = (
                max(probabilities)
                * 100
            )

        else:

            confidence = 0


        return jsonify({

            "success": True,

            "filename":
                file.filename,

            "extracted_text":
                extracted_text,

            "language":
                str(prediction),

            "confidence":
                round(
                    confidence,
                    2
                )

        })


    except Exception as error:

        print(
            "File error:",
            error
        )


        return jsonify({

            "success": False,

            "message":
                str(error)

        }), 500


# =========================================================
# TEXT CHUNKING
# =========================================================

def split_text_into_chunks(
    text,
    max_chars=400
):

    words = text.split()

    chunks = []

    current = ""


    for word in words:

        test = (
            current
            + " "
            + word
        ).strip()


        if len(test) <= max_chars:

            current = test

        else:

            if current:

                chunks.append(
                    current
                )

            current = word


    if current:

        chunks.append(
            current
        )


    return chunks


# =========================================================
# TRANSLATION CHUNK
# =========================================================

def translate_chunk(
    text,
    source_code,
    target_code
):

    # Same language

    if source_code == target_code:

        return text


    try:

        # MyMemory Translation API

        url = (
            "https://api.mymemory.translated.net/get"
        )


        params = {

            "q":
                text,

            "langpair":
                f"{source_code}|{target_code}"

        }


        response = requests.get(

            url,

            params=params,

            timeout=30

        )


        if response.status_code != 200:

            raise Exception(

                "Translation API error: "
                + str(
                    response.status_code
                )

            )


        data = response.json()


        if data.get(
            "responseStatus"
        ) != 200:

            raise Exception(

                data.get(

                    "responseDetails",

                    "Translation failed."

                )

            )


        translated = (

            data

            .get(
                "responseData",
                {}
            )

            .get(
                "translatedText",
                ""
            )

        )


        if not translated:

            raise Exception(
                "Empty translation received."
            )


        return translated


    except Exception as error:

        print(
            "Translation provider error:",
            error
        )


        raise Exception(

            "Translation service unavailable."

        )


# =========================================================
# TRANSLATE TEXT
# =========================================================

@app.route(
    "/translate",
    methods=["POST"]
)
def translate():

    try:

        data = request.get_json()


        if not data:

            return jsonify({

                "success": False,

                "message":
                    "No JSON data received."

            }), 400


        text = data.get(
            "text",
            ""
        ).strip()


        source_language = data.get(

            "source_language",

            "English"

        )


        target_language = data.get(

            "target_language",

            "English"

        )


        if not text:

            return jsonify({

                "success": False,

                "message":
                    "Please provide text."

            }), 400


        # Get language codes

        source_code = (
            LANGUAGE_CODES.get(
                source_language
            )
        )


        target_code = (
            LANGUAGE_CODES.get(
                target_language
            )
        )


        if not source_code:

            return jsonify({

                "success": False,

                "message":
                    "Unsupported source language."

            }), 400


        if not target_code:

            return jsonify({

                "success": False,

                "message":
                    "Unsupported target language."

            }), 400


        # Same language

        if source_code == target_code:

            translated_text = text

            total_chunks = 1


        else:

            chunks = (
                split_text_into_chunks(
                    text,
                    400
                )
            )


            translated_chunks = []

            total_chunks = len(
                chunks
            )


            for index, chunk in enumerate(

                chunks,

                start=1

            ):

                print(

                    f"Translating chunk "
                    f"{index}/{total_chunks}"

                )


                translated = (
                    translate_chunk(

                        chunk,

                        source_code,

                        target_code

                    )
                )


                translated_chunks.append(
                    translated
                )


                # Small delay

                time.sleep(
                    1.2
                )


            translated_text = (
                "\n\n".join(
                    translated_chunks
                )
            )


        return jsonify({

            "success": True,

            "original_text":
                text,

            "translated_text":
                translated_text,

            "source_language":
                source_language,

            "target_language":
                target_language,

            "chunks":
                total_chunks

        })


    except Exception as error:

        print(
            "Translation error:",
            error
        )


        return jsonify({

            "success": False,

            "message":
                str(error)

        }), 500


# =========================================================
# DOWNLOAD TRANSLATED FILE
# =========================================================

@app.route(
    "/download",
    methods=["POST"]
)
def download_document():

    try:

        data = request.get_json()


        if not data:

            return jsonify({

                "success": False,

                "message":
                    "No JSON data received."

            }), 400


        text = data.get(
            "text",
            ""
        ).strip()


        file_format = data.get(
            "format",
            "txt"
        ).lower()


        if not text:

            return jsonify({

                "success": False,

                "message":
                    "No translated text available."

            }), 400


        # =================================================
        # TXT
        # =================================================

        if file_format == "txt":

            output = BytesIO()


            output.write(
                text.encode(
                    "utf-8"
                )
            )


            output.seek(0)


            return send_file(

                output,

                as_attachment=True,

                download_name=
                    "translated_text.txt",

                mimetype=
                    "text/plain"

            )


        # =================================================
        # DOCX
        # =================================================

        elif file_format == "docx":

            document = Document()


            for paragraph in (
                text.split("\n")
            ):

                document.add_paragraph(
                    paragraph
                )


            output = BytesIO()


            document.save(
                output
            )


            output.seek(0)


            return send_file(

                output,

                as_attachment=True,

                download_name=
                    "translated_document.docx",

                mimetype=
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

            )


        # =================================================
        # PDF
        # =================================================

        elif file_format == "pdf":

            output = BytesIO()


            pdf = canvas.Canvas(

                output,

                pagesize=A4

            )


            width, height = A4


            left_margin = 45

            top_margin = (
                height - 50
            )

            y = top_margin

            line_height = 16


            for paragraph in (
                text.split("\n")
            ):

                words = paragraph.split()

                line = ""


                for word in words:

                    test_line = (

                        line
                        + " "
                        + word

                    ).strip()


                    if len(test_line) > 85:

                        pdf.drawString(

                            left_margin,

                            y,

                            line

                        )


                        y -= line_height

                        line = word


                        if y < 50:

                            pdf.showPage()

                            y = top_margin


                    else:

                        line = test_line


                if line:

                    pdf.drawString(

                        left_margin,

                        y,

                        line

                    )

                    y -= line_height


                y -= 8


                if y < 50:

                    pdf.showPage()

                    y = top_margin


            pdf.save()


            output.seek(0)


            return send_file(

                output,

                as_attachment=True,

                download_name=
                    "translated_document.pdf",

                mimetype=
                    "application/pdf"

            )


        # =================================================
        # INVALID FORMAT
        # =================================================

        else:

            return jsonify({

                "success": False,

                "message":
                    "Unsupported download format."

            }), 400


    except Exception as error:

        print(
            "Download error:",
            error
        )


        return jsonify({

            "success": False,

            "message":
                str(error)

        }), 500


# =========================================================
# RUN SERVER
# =========================================================

if __name__ == "__main__":

    print(
        "======================================"
    )

    print(
        " LANGUAGE DETECTION FULL STACK APP"
    )

    print(
        "======================================"
    )

    print(
        "Frontend + Backend running together"
    )

    print(
        "======================================"
    )


    app.run(

        host="0.0.0.0",

        port=5000,

        debug=True

    )